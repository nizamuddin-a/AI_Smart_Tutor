
import re
from pathlib import Path

import fitz
from docx import Document as WordDocument
from pptx import Presentation


def extract_document(path: str, extension: str) -> list[dict]:
    """Return text blocks with page/slide numbers."""
    extension = extension.lower()
    pages = []

    if extension == ".pdf":
        pdf = fitz.open(path)
        for page_number, page in enumerate(pdf, start=1):
            pages.append({
                "page": page_number,
                "text": page.get_text("text")
            })
        pdf.close()

    elif extension == ".pptx":
        presentation = Presentation(path)
        for slide_number, slide in enumerate(presentation.slides, start=1):
            lines = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    if shape.text.strip():
                        lines.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        lines.append(" | ".join(
                            cell.text.strip() for cell in row.cells
                        ))
            pages.append({
                "page": slide_number,
                "text": "\n".join(lines)
            })

    elif extension == ".docx":
        document = WordDocument(path)
        text = "\n".join(
            paragraph.text for paragraph in document.paragraphs
            if paragraph.text.strip()
        )
        pages.append({"page": 1, "text": text})

    elif extension in {".txt", ".md"}:
        text = Path(path).read_text(encoding="utf-8", errors="replace")
        pages.append({"page": 1, "text": text})

    else:
        raise ValueError("Supported formats: PDF, PPTX, DOCX, TXT, MD.")

    return pages


def make_chunks(pages: list[dict], chunk_size: int = 1000,
                overlap: int = 150) -> list[dict]:
    chunks = []
    chunk_id = 0

    for page in pages:
        text = re.sub(r"\s+", " ", page["text"]).strip()
        if not text:
            continue

        start = 0
        while start < len(text):
            end = min(start + chunk_size, len(text))
            piece = text[start:end].strip()

            if piece:
                chunks.append({
                    "chunk_id": chunk_id,
                    "page": page["page"],
                    "text": piece
                })
                chunk_id += 1

            if end >= len(text):
                break
            start = max(end - overlap, start + 1)

    return chunks


def retrieve_chunks(chunks: list[dict], question: str,
                    limit: int = 5) -> list[dict]:
    """Simple local keyword retrieval; no embedding API needed."""
    stop_words = {
        "the", "and", "for", "that", "this", "with", "what",
        "how", "are", "was", "can", "from", "into", "about",
        "explain", "please", "tell", "me", "you", "does"
    }
    terms = {
        word for word in re.findall(r"[a-zA-Z0-9]+", question.lower())
        if len(word) > 2 and word not in stop_words
    }

    if not terms:
        return chunks[:limit]

    scored = []
    for chunk in chunks:
        text = chunk["text"].lower()
        score = sum(text.count(term) for term in terms)
        if score > 0:
            scored.append((score, chunk))

    scored.sort(key=lambda item: item[0], reverse=True)
    return [chunk for _, chunk in scored[:limit]]